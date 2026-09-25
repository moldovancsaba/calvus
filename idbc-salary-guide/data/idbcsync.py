#!/usr/bin/env python3
"""Build the IDBC Salary Guide from the IDBCSYNC tab of the IDBC_bertabla Google Sheet.

IDBCSYNC is the single source of every value on the guide: page texts, area texts, the SAP
catalogue, salary rows, the Expert Pool and the survey results. Each row is one value
(`id`, `változó`, `érték`, `megjelenés`, `segítség`); only `id` and `érték` are read here.

    python3 idbc-salary-guide/data/idbcsync.py <IDBC_bertabla.xlsx>           # build
    python3 idbc-salary-guide/data/idbcsync.py <IDBC_bertabla.xlsx> --check   # validate only

Writes data/guide-data.json, data/areas.json, and the texts into the pages (every element
marked data-sync / data-sync-attr / data-sync-href, every script literal marked /*sync:ID*/).
Only files whose content changes are written. Any error stops the build with exit 1 and
writes nothing, so the live site keeps the last good version (docs/12-technical-design.md).
"""
import datetime, html, json, pathlib, re, sys

ROOT = pathlib.Path(__file__).resolve().parent.parent          # idbc-salary-guide/
PAGES = ["index.html", "berezes/index.html", "sap/index.html", "expert-pool/index.html",
         "esettanulmanyok/index.html", "regisztracio/index.html", "kapcsolat/index.html",
         "terulet/index.html"]
CHART = "assets/top3-chart.js"
SHEET = "IDBCSYNC"
TOTAL_KEY = "__total__"
SIDES = ("employee", "employer")
# Placeholders a script text may contain; anything else in {braces} is a typo and stops the build.
TOKENS = {"HOME-JS-TOOLTIP-COUNT": {"n"}, "SHARED-JS-SCALE-HINT": {"max"},
          "TERULET-JS-TITLE": {"terulet"}, "TERULET-JS-VIDEO-ARIA": {"terulet"},
          "TERULET-JS-KEYTHOUGHT": {"terulet"}, "BEREK-JS-TOP3-TITLE": {"terulet"},
          "BEREK-JS-CHART-LABEL": {"terulet"}, "POOL-JS-TIP-MARKET": {"n"}, "POOL-JS-MARKET": {"n"},
          "POOL-JS-TIP": {"iparag", "pozicio", "darab", "osszes", "arany"},
          "SHARED-CHART-TALENT-PILL": {"n"}}
MEDIA = {"video", "highlight"}
YES, NO = {"igen", "i", "x", "true", "1", "yes"}, {"nem", "n", "", "false", "0", "no"}

errors, warnings = [], []


# ---------------------------------------------------------------- reading the sheet
def cell_text(v, rid):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "igen" if v else "nem"
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        errors.append(f"{rid}: a Sheets dátummá alakította az értéket ({v}) — írd be újra "
                      f"aposztróffal az elején (pl. '1-5), vagy állítsd a cellát egyszerű szövegre")
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v).replace("\r\n", "\n").strip()


def read_sheet(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"no {SHEET} tab in {path} (tabs: {', '.join(wb.sheetnames)})")
    rows, head = [], None
    for r in wb[SHEET].iter_rows(values_only=True):
        if head is None:
            names = [str(c or "").strip().lower() for c in r]
            if "id" in names and "érték" in names:
                head = (names.index("id"), names.index("érték"))
            continue
        rid = str(r[head[0]] or "").strip() if len(r) > head[0] else ""
        if not rid or rid.startswith("#"):
            continue
        rows.append((rid, cell_text(r[head[1]] if len(r) > head[1] else None, rid)))
    if head is None:
        sys.exit(f"{SHEET}: no header row with 'id' and 'érték'")
    values = {}
    for rid, v in rows:
        if rid in values:
            errors.append(f"{rid}: az azonosító kétszer szerepel")
        values[rid] = v
    return values, [rid for rid, _ in rows]


# ---------------------------------------------------------------- value parsers
def as_int(rid, v, required=False):
    s = re.sub(r"[\s .]", "", v).replace("Ft", "")
    if s == "":
        if required:
            errors.append(f"{rid}: üres, számot várunk")
        return None
    if not re.fullmatch(r"\d+", s):
        errors.append(f"{rid}: „{v}” nem egész szám")
        return None
    return int(s)


def as_float(rid, v, lo=None, hi=None):
    s = v.replace(",", ".").replace("%", "").strip()
    try:
        f = float(s)
    except ValueError:
        errors.append(f"{rid}: „{v}” nem szám")
        return 0.0
    if (lo is not None and f < lo) or (hi is not None and f > hi):
        errors.append(f"{rid}: {v} kívül esik a {lo}–{hi} tartományon")
    return f


def as_bool(rid, v):
    s = v.strip().lower()
    if s in YES:
        return True
    if s not in NO:
        errors.append(f"{rid}: „{v}” — igen vagy nem")
    return False


def num_out(f):
    return f


# ---------------------------------------------------------------- the model
def build(values, order):
    first = {rid: i for i, rid in enumerate(order)}

    def grouped(pattern):
        """{group key: {field: value}} for ids matching pattern (groups: key, field), in sheet order."""
        out = {}
        for rid in order:
            m = re.fullmatch(pattern, rid)
            if m:
                out.setdefault(m.group(1), {})[m.group(2)] = (rid, values[rid])
        return out

    # --- areas (Piaci trendek) ---
    areas = []
    for key, f in grouped(r"AREA-(.+)-(NAME|SUMMARY|MEDIA|EDITION)").items():
        get = lambda k: f.get(k, (f"AREA-{key}-{k}", ""))
        rid, name = get("NAME")
        if not name:
            errors.append(f"{rid}: a terület neve üres")
        media = get("MEDIA")[1].strip().lower()
        if media not in MEDIA:
            errors.append(f"{get('MEDIA')[0]}: „{media}” — video vagy highlight")
        summary = get("SUMMARY")[1]
        if not summary:
            errors.append(f"{get('SUMMARY')[0]}: az összefoglaló üres")
        areas.append({"slug": key.lower(), "name": name, "edition": get("EDITION")[1],
                      "summary": summary, "media": media})

    # --- SAP catalogue ---
    sap = []
    for key, f in grouped(r"SAPPROD-(CAT\d+)-(NAME|ITEM\d+)").items():
        items = [v for k, (rid, v) in sorted(f.items(), key=lambda kv: first[kv[1][0]]) if k != "NAME" and v]
        name = f.get("NAME", ("", ""))[1]
        if items and not name:
            errors.append(f"SAPPROD-{key}-NAME: a kategória neve üres")
        if name:
            sap.append({"category": name, "items": items})

    # --- salary (Bérek + SAP) ---
    sal_names = {m.group(1): (rid, values[rid]) for rid in order
                 for m in [re.fullmatch(r"SAL-(.+)-NAME", rid)] if m}
    salary = []
    for key, f in grouped(r"SAL-(.+-\d{3})-(POZICIO|SZINT|TOP3|MIN|IDBC|MAX|JUTTATAS|LINKEDIN)").items():
        g = lambda k: f.get(k, (f"SAL-{key}-{k}", ""))
        if not g("POZICIO")[1]:
            continue                                   # a cleared position is a deleted row
        kod = key.rsplit("-", 1)[0]
        if kod not in sal_names:
            errors.append(f"SAL-{key}: nincs SAL-{kod}-NAME sor a területnévhez")
            continue
        row = {"id": key, "kod": kod, "terulet": sal_names[kod][1], "szint": g("SZINT")[1] or None,
               "pozicio": g("POZICIO")[1], "top3": as_bool(g("TOP3")[0], g("TOP3")[1]),
               "min": as_int(g("MIN")[0], g("MIN")[1]), "idbc": as_int(g("IDBC")[0], g("IDBC")[1]),
               "max": as_int(g("MAX")[0], g("MAX")[1]), "juttatas": g("JUTTATAS")[1]}
        li = as_int(g("LINKEDIN")[0], g("LINKEDIN")[1])
        if li is not None:
            row["linkedin"] = li
        salary.append(row)

    # --- Expert Pool ---
    pool = []
    for key, f in grouped(r"EXPERT-(\d+)-(IPARAG|POZICIO|DARAB|LINKEDIN)").items():
        g = lambda k: f.get(k, (f"EXPERT-{key}-{k}", ""))
        if not g("POZICIO")[1]:
            continue
        darab = as_int(g("DARAB")[0], g("DARAB")[1])
        if darab is None:
            warnings.append(f"EXPERT-{key} ({g('POZICIO')[1]}): nincs darabszám — a csempe addig nem jelenik meg")
            continue
        row = {"iparag": g("IPARAG")[1], "pozicio": g("POZICIO")[1], "darab": darab}
        li = as_int(g("LINKEDIN")[0], g("LINKEDIN")[1])
        if li is not None:
            row["linkedin"] = li
        pool.append(row)

    # --- survey: definitions ---
    total_label = values.get("SURVEY-TOTAL-LABEL", "")
    if not total_label:
        errors.append("SURVEY-TOTAL-LABEL: üres")
    sets = {k: v for k, v in ((m.group(1), values[rid]) for rid in order
                              for m in [re.fullmatch(r"SURVEY-(SET\d+)-NAME", rid)] if m)}
    segments = {side: [values[rid] for rid in order if re.fullmatch(rf"SURVEY-{side.upper()}-SEG\d+", rid)]
                for side in SIDES}
    questions = {}                                     # (side, "Qnn") -> definition
    for rid in order:
        m = re.fullmatch(r"SURVEY-(EMPLOYEE|EMPLOYER)-(Q\d+)-(TEXT|OPT\d+|SCALEMAX|SORREND)", rid)
        if not m:
            continue
        q = questions.setdefault((m.group(1).lower(), m.group(2)), {"opts": {}, "text": "", "scale": None, "keep": False})
        if m.group(3) == "TEXT":
            q["text"] = values[rid]
            if not values[rid]:
                errors.append(f"{rid}: a kérdés szövege üres")
        elif m.group(3) == "SCALEMAX":
            q["scale"] = as_int(rid, values[rid], required=True)
        elif m.group(3) == "SORREND":
            q["keep"] = as_bool(rid, values[rid])
        else:
            q["opts"][m.group(3)] = values[rid]
    texts = {}
    for (side, qn), q in questions.items():
        if q["text"] in texts.get(side, {}):
            errors.append(f"SURVEY-{side.upper()}-{qn}-TEXT: ugyanez a kérdésszöveg máshol is szerepel ezen az oldalon")
        texts.setdefault(side, {})[q["text"]] = qn
    topics = []
    for key, f in grouped(r"SURVEY-(TOPIC\d+)-(NAME)").items():
        topic = {"topic": f["NAME"][1], "questionSets": {}}
        for rid in order:
            m = re.fullmatch(rf"SURVEY-{key}-(SET\d+)-(EMPLOYEE|EMPLOYER)-\d+", rid)
            if not m:
                continue
            ref = values[rid]
            side = m.group(2).lower()
            q = questions.get((side, ref))
            if q is None:
                errors.append(f"{rid}: „{ref}” nem létező kérdés")
                continue
            label = {"employee": "Munkavállalók", "employer": "Munkáltatók"}[side]
            ed = topic["questionSets"].setdefault(sets.get(m.group(1), m.group(1)), {"Munkavállalók": [], "Munkáltatók": []})
            ed[label].append(q["text"])
        topics.append(topic)

    # --- survey: values per dataset ---
    ds_order = ["TOTAL"] + [a["slug"].upper() for a in areas]
    raw = {}
    for rid in order:
        m = re.fullmatch(r"SURVEY-(.+)-(EMPLOYEE|EMPLOYER)-(Q\d+)-(?:(SEG\d+)-)?(N|OPT\d+)", rid)
        if m:
            ds, side, qn, seg, what = m.groups()
            raw.setdefault(ds, {}).setdefault(side.lower(), {}).setdefault(qn, {}).setdefault(seg, {})[what] = (rid, values[rid])
    datasets = {}
    edition_of = {a["slug"].upper(): a["edition"] for a in areas}
    edition_of["TOTAL"] = values.get("SURVEY-TOTAL-EDITION", "")
    label_of = {a["slug"].upper(): a["name"] for a in areas}
    label_of["TOTAL"] = total_label
    for ds in ds_order:
        if ds not in raw:
            continue
        if edition_of[ds] not in sets.values():
            errors.append(f"{'SURVEY-TOTAL-EDITION' if ds == 'TOTAL' else f'AREA-{ds}-EDITION'}: "
                          f"„{edition_of[ds]}” nem kérdéssor ({', '.join(sets.values())})")
        out = {"label": label_of[ds], "questionSet": edition_of[ds]}
        for side in SIDES:
            base, cross = {}, {}
            for qn in sorted(raw[ds].get(side, {}), key=lambda k: int(k[1:])):
                q = questions.get((side, qn))
                if q is None:
                    errors.append(f"SURVEY-{ds}-{side.upper()}-{qn}: nincs ilyen kérdés (SURVEY-{side.upper()}-{qn}-TEXT)")
                    continue
                opt_keys = sorted(q["opts"], key=lambda k: int(k[3:]))
                cells = raw[ds][side][qn]
                if None in cells:                      # whole-sample answers
                    c = cells[None]
                    if q["scale"]:
                        items = [{"label": q["opts"][o], "value": num_out(as_float(c[o][0], c[o][1], 0, q["scale"]))}
                                 for o in opt_keys if o in c]
                        base[q["text"]] = {"kind": "weighted", "scaleMax": q["scale"], "items": items}
                    else:
                        n = as_int(c["N"][0], c["N"][1], required=True) if "N" in c else None
                        items = []
                        for o in opt_keys:
                            if o in c:
                                pct = as_float(c[o][0], c[o][1], 0, 100)
                                items.append({"label": q["opts"][o], "percent": num_out(pct),
                                              "count": round(pct * (n or 0) / 100)})
                        base[q["text"]] = {"kind": "simple", "items": items}
                segs = sorted((s for s in cells if s), key=lambda s: int(s[3:]))
                if segs:
                    groups = [o for o in opt_keys if any(o in cells[s] for s in segs)]
                    items = []
                    for s in segs:
                        c = cells[s]
                        si = int(s[3:]) - 1
                        if si >= len(segments[side]):
                            errors.append(f"SURVEY-{ds}-{side.upper()}-{qn}-{s}: nincs SURVEY-{side.upper()}-{s} szegmens")
                            continue
                        n = as_int(c["N"][0], c["N"][1], required=True) if "N" in c else 0
                        vals = []
                        for o in groups:
                            pct = as_float(c[o][0], c[o][1], 0, 100) if o in c else 0.0
                            vals.append({"group": q["opts"][o], "percent": num_out(pct), "count": round(pct * (n or 0) / 100)})
                        items.append({"label": segments[side][si], "values": vals})
                    cross[q["text"]] = {"kind": "matrix-percent", "groups": [q["opts"][o] for o in groups], "items": items}
            out[side] = {"base": base, "cross": cross}
        datasets[TOTAL_KEY if ds == "TOTAL" else ds.lower()] = out

    keep = [q["text"] for q in questions.values() if q["keep"]]
    guide = {"totalKey": TOTAL_KEY, "totalLabel": total_label, "segments": segments,
             "keepSourceOrder": keep, "topics": topics, "datasets": datasets,
             "salary": {"webBertabla": salary, "expertPool": pool}, "sapProducts": sap}
    return guide, {"areas": areas}


# ---------------------------------------------------------------- page texts
BR = re.compile(r"<br\s*/?>", re.I)
WS = re.compile(r"[ \t\r\n]+")
START_TAG = re.compile(r'<([a-zA-Z][\w-]*)((?:\s+[^\s=>/]+(?:\s*=\s*"[^"]*")?)*)\s*/?>')
JS_TEXT = re.compile(r'/\*sync:([A-Z0-9-]+)\*/"((?:[^"\\\n]|\\.)*)"')


def text_value(inner):
    return "\n".join(WS.sub(" ", html.unescape(p)).strip() for p in BR.split(inner)).strip("\n")


def text_html(value):
    return "<br>".join(html.escape(line, quote=False) for line in value.split("\n"))


def need(values, rid, where):
    if rid not in values:
        errors.append(f"{rid}: hiányzik a táblázatból (használja: {where})")
        return None
    return values[rid]


def sync_html(src, values, where):
    out, pos = [], 0
    for m in START_TAG.finditer(src):
        attrs = m.group(2)
        if "data-sync" not in attrs:
            continue
        tag, start_tag = m.group(1).lower(), m.group(0)
        new_tag = start_tag
        spec = re.search(r'\sdata-sync-attr="([^"]+)"', attrs)
        for part in (spec.group(1).split(";") if spec else []):
            name, rid = part.split(":", 1)
            val = need(values, rid, where)
            am = re.search(rf'\s{re.escape(name)}="([^"]*)"', new_tag)
            if val is not None and am and WS.sub(" ", html.unescape(am.group(1))).strip() != val:
                new_tag = new_tag[:am.start(1)] + html.escape(val, quote=True) + new_tag[am.end(1):]
        hm = re.search(r'\sdata-sync-href="(tel|mailto):([^"]+)"', attrs)
        if hm:
            val = need(values, hm.group(2), where)
            href = re.search(r'\shref="([^"]*)"', new_tag)
            if val is not None and href:
                addr = re.sub(r"[^\d+]", "", val) if hm.group(1) == "tel" else val.strip()
                query = href.group(1).split("?", 1)[1] if "?" in href.group(1) else None
                want = f"{hm.group(1)}:{addr}" + (f"?{query}" if query else "")
                if href.group(1) != want:
                    new_tag = new_tag[:href.start(1)] + html.escape(want, quote=True) + new_tag[href.end(1):]
        out.append(src[pos:m.start()])
        out.append(new_tag)
        pos = m.end()
        tm = re.search(r'\sdata-sync="([^"]+)"', attrs)
        if tm:
            val = need(values, tm.group(1), where)
            close = src.index(f"</{tag}>", m.end())
            inner = src[m.end():close]
            if "<" in BR.sub("", re.sub(r"<!--.*?-->", "", inner, flags=re.S)):
                errors.append(f"{tm.group(1)}: {where} — a jelölt elem nem csak szöveget tartalmaz")
            elif val is not None and text_value(inner) != val:
                out.append(text_html(val))
                pos = close
    out.append(src[pos:])
    return "".join(out)


def sync_js(src, values, where):
    def rep(m):
        rid = m.group(1)
        val = need(values, rid, where)
        if val is None:
            return m.group(0)
        bad = set(re.findall(r"\{(\w+)\}", val)) - TOKENS.get(rid, set())
        if bad:
            errors.append(f"{rid}: ismeretlen helyőrző: {', '.join('{' + b + '}' for b in sorted(bad))}"
                          f" (használható: {', '.join('{' + t + '}' for t in sorted(TOKENS.get(rid, ()))) or 'nincs'})")
            return m.group(0)
        if json.loads(f'"{m.group(2)}"') == val:
            return m.group(0)
        return f"/*sync:{rid}*/" + json.dumps(val, ensure_ascii=False).replace("</", "<\\/")
    return JS_TEXT.sub(rep, src)


def sync_page(src, values, where):
    parts, pos = [], 0
    for m in re.finditer(r"(<script\b[^>]*>)(.*?)(</script>)", src, re.S):
        parts.append(sync_html(src[pos:m.start()], values, where))
        parts.append(m.group(1) + sync_js(m.group(2), values, where) + m.group(3))
        pos = m.end()
    parts.append(sync_html(src[pos:], values, where))
    return "".join(parts)


def referenced_ids(text):
    ids = set(re.findall(r'data-sync="([^"]+)"', text)) | set(JS_TEXT_ID.findall(text))
    for spec in re.findall(r'data-sync-attr="([^"]+)"', text):
        ids |= {p.split(":", 1)[1] for p in spec.split(";")}
    ids |= {m for m in re.findall(r'data-sync-href="(?:tel|mailto):([^"]+)"', text)}
    return ids


JS_TEXT_ID = re.compile(r"/\*sync:([A-Z0-9-]+)\*/")


# ---------------------------------------------------------------- main
def dump(obj, compact=False):
    if compact:
        return json.dumps(obj, ensure_ascii=False, separators=(",", ":"))
    return json.dumps(obj, ensure_ascii=False, indent=2) + "\n"


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 1:
        sys.exit(__doc__)
    check_only = "--check" in sys.argv
    values, order = read_sheet(args[0])
    guide, areas = build(values, order)

    outputs = {ROOT / "data/guide-data.json": dump(guide, compact=True), ROOT / "data/areas.json": dump(areas)}
    used = set()
    for rel in PAGES + [CHART]:
        p = ROOT / rel
        src = p.read_text(encoding="utf-8")
        used |= referenced_ids(src)
        outputs[p] = sync_js(src, values, rel) if rel.endswith(".js") else sync_page(src, values, rel)
    chart = ROOT / CHART
    if outputs[chart] != chart.read_text(encoding="utf-8"):
        # The chart script is cached by ?v=; a changed text needs a new version on every page.
        vers = [int(v) for rel in PAGES for v in re.findall(r"top3-chart\.js\?v=(\d+)", outputs[ROOT / rel])]
        nxt = max(vers) + 1
        for rel in PAGES:
            outputs[ROOT / rel] = re.sub(r"top3-chart\.js\?v=\d+", f"top3-chart.js?v={nxt}", outputs[ROOT / rel])

    data_ids = re.compile(r"(AREA|SAPPROD|SAL|EXPERT|SURVEY)-")
    unused = [rid for rid in order if rid not in used and not data_ids.match(rid)]
    if unused:
        warnings.append(f"{len(unused)} sor egyik oldalon sem szerepel: {', '.join(unused[:8])}{' …' if len(unused) > 8 else ''}")

    for w in warnings:
        print("figyelem:", w)
    if errors:
        for e in dict.fromkeys(errors):
            print("HIBA:", e)
        print(f"{len(errors)} hiba — semmi nem íródott ki, az oldal marad a legutóbbi jó változaton.")
        sys.exit(1)
    changed = [p for p, text in outputs.items() if p.read_text(encoding="utf-8") != text]
    print(f"{SHEET}: {len(values)} sor; {len(areas['areas'])} terület, {len(guide['salary']['webBertabla'])} bérsor, "
          f"{len(guide['salary']['expertPool'])} Expert Pool sor, {len(guide['datasets'])} kutatási adatkör")
    if check_only:
        print("ellenőrzés: rendben;", len(changed), "fájl változna:", ", ".join(str(p.relative_to(ROOT)) for p in changed) or "egy sem")
        return
    for p in changed:
        p.write_text(outputs[p], encoding="utf-8")
    print(len(changed), "fájl frissült:", ", ".join(str(p.relative_to(ROOT)) for p in changed) or "egy sem")


if __name__ == "__main__":
    main()
