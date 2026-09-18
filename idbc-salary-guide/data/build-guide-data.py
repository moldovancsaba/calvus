# -*- coding: utf-8 -*-
"""Rebuild the survey half of guide-data.json from the client's final Excel.

    python3 idbc-salary-guide/data/build-guide-data.py <Salary Guide- kutatási eredmények-final.xlsx>

The salary half (webBertabla, expertPool, sapProducts, ...) is preserved untouched —
the workbook only carries survey results; build-salary-data.py owns that half.
"""
import io, json, sys
import openpyxl

if len(sys.argv) != 2:
    sys.exit("usage: build-guide-data.py <research workbook .xlsx>")
SRC = sys.argv[1]
OUT = __file__.rsplit("/", 1)[0] + "/guide-data.json"

wb = openpyxl.load_workbook(SRC, data_only=True)

# Sheet prefix -> the area name used by the site (data/areas.json "name").
AREA_SHEETS = {
    "IT": "IT",
    "IT Contracting": "IT Contracting",
    "Finance": "Pénzügy és számvitel",
    "Sales & Marketing": "Sales & Marketing",
    "HR": "HR",
    "BSC": "BSC",
    "Ép, ingatlan": "Építőipar, ingatlan",
    "Gyártás": "Gyártás, termelés, mérnökség",
    "Üsz, Admin": "Office Support & Ügyfélszolgálat",
    "Logisztika": "Logisztika és szállítás",
    "Pharma": "Pharma & Life Sciences",
}
# Which question set each dataset follows.
IT_SET = {"IT", "IT Contracting"}
TOTAL_KEY = "__total__"
TOTAL_LABEL = "Összesített adatok"

# The two questions the client wants reduced to the weighted average only.
WEIGHTED_ONLY = (
    "Jelöld 1-4-ig terjedő skálán, hogy mennyire fontosak",
)

import re
def cell(ws, r, c):
    """Header and question cells carry hard line breaks; collapse all whitespace so the
    same question text is one key everywhere (topic sheet, totals, crosstabs)."""
    v = ws.cell(r, c).value
    return re.sub(r"\s+", " ", v).strip() if isinstance(v, str) else v

def pct(v):
    """Workbook stores 0-1 fractions; the site works in percentages."""
    return round(float(v) * 100, 4) if isinstance(v, (int, float)) else 0.0

def is_weighted_only(q):
    return any(q.startswith(p) for p in WEIGHTED_ONLY)

def parse_total(ws):
    """A 'Total' sheet: a sequence of question blocks keyed by the question text."""
    out, r, n = {}, 1, ws.max_row
    while r <= n:
        title = cell(ws, r, 1)
        if not isinstance(title, str) or not title or cell(ws, r, 2) is not None:
            r += 1
            continue
        # find this block's header row
        hr = None
        for rr in range(r + 1, min(r + 5, n) + 1):
            if cell(ws, rr, 2) == "Response Percent":
                hr = rr
                break
        if hr is None:
            r += 1
            continue
        weighted_col = 10 if cell(ws, hr, 10) == "Súlyozott átlag" else None
        groups = None
        if weighted_col:  # scale matrix: group labels sit one row above the header
            groups = [str(cell(ws, hr - 1, c)) for c in (2, 4, 6, 8)]
        items, rr = [], hr + 1
        while rr <= n:
            label = cell(ws, rr, 1)
            if not isinstance(label, str) or not label:
                break
            if weighted_col:
                wv = cell(ws, rr, weighted_col)
                items.append({
                    "label": label,
                    "value": round(float(wv), 3) if isinstance(wv, (int, float)) else 0.0,
                    "values": [{"group": groups[i], "percent": pct(cell(ws, rr, 2 + i * 2)),
                                "count": cell(ws, rr, 3 + i * 2) or 0} for i in range(4)],
                })
            else:
                items.append({"label": label, "percent": pct(cell(ws, rr, 2)),
                              "count": cell(ws, rr, 3) or 0})
            rr += 1
        if items:
            if weighted_col:
                # Client (2026-09-08): show only the weighted average for these questions.
                out[title] = {"kind": "weighted", "scaleMax": 4,
                              "items": [{"label": i["label"], "value": i["value"]} for i in items]}
            else:
                out[title] = {"kind": "simple", "items": items}
        r = rr
    return out

def parse_cross(ws):
    """A crosstab sheet: blocks titled '<dimension> (vs) <question>'."""
    out, r, n = {}, 1, ws.max_row
    while r <= n:
        title = cell(ws, r, 1)
        if not isinstance(title, str) or "(vs)" not in (title or ""):
            r += 1
            continue
        question = title.split("(vs)", 1)[1].strip()
        # answer labels row, then the Response Percent/Count header row
        lr = None
        for rr in range(r + 1, min(r + 5, n) + 1):
            if cell(ws, rr, 2) == "Response Percent":
                lr = rr
                break
        if lr is None:
            r += 1
            continue
        groups, c = [], 2
        while c <= ws.max_column:
            g = cell(ws, lr - 1, c)
            if isinstance(g, str) and g:
                groups.append(g)
                c += 2
            else:
                break
        items, rr = [], lr + 1
        while rr <= n:
            label = cell(ws, rr, 1)
            if not isinstance(label, str) or not label:
                break
            items.append({"label": label, "values": [
                {"group": groups[i], "percent": pct(cell(ws, rr, 2 + i * 2)),
                 "count": cell(ws, rr, 3 + i * 2) or 0} for i in range(len(groups))]})
            rr += 1
        if items and groups:
            out[question] = {"kind": "matrix-percent", "groups": groups, "items": items}
        r = rr
    return out

def parse_topics():
    """'Téma besorolás': topic -> question list, per side, for each question set."""
    ws = wb["Téma besorolás"]
    sets = {"Általános": 1, "IT + Contracting": 5}   # first column of each block
    topics_order, data = [], {}
    for set_name, c0 in sets.items():
        current = None
        for r in range(1, ws.max_row + 1):
            a, b = cell(ws, r, c0), cell(ws, r, c0 + 1)
            if not a:
                continue
            if a == set_name:
                continue
            if a in ("Munkavállalók", "Munkáltatók"):
                continue
            # a topic header is a lone cell whose next row is the side header
            nxt = cell(ws, r + 1, c0)
            if nxt == "Munkavállalók" and not b:
                current = a
                if current not in topics_order:
                    topics_order.append(current)
                data.setdefault(current, {}).setdefault(set_name, {"Munkavállalók": [], "Munkáltatók": []})
                continue
            if current:
                d = data[current][set_name]
                if a:
                    d["Munkavállalók"].append(a)
                if b:
                    d["Munkáltatók"].append(b)
    return topics_order, data

# ---- build ----
topics_order, topic_data = parse_topics()

def find_sheet(prefix, side, kind):
    """Excel truncates sheet names at 31 chars, so match by prefix instead of exact name."""
    want_pref = "" if prefix is None else prefix + " "
    cands = []
    for n in wb.sheetnames:
        if prefix is None:
            if not n.startswith(side + " "):
                continue
        elif not n.startswith(want_pref + side + " "):
            continue
        tail = n[len(want_pref) + len(side) + 1:]
        if kind == "total" and tail.startswith("Total"):
            cands.append(n)
        if kind == "cross" and (tail.startswith("Tapaszt") or tail.startswith("Cégmé")):
            cands.append(n)
    if len(cands) != 1:
        sys.exit(f"sheet lookup failed: prefix={prefix!r} side={side} kind={kind} -> {cands}")
    return cands[0]

def sheets_for(prefix):
    """Returns (b2c_total, b2c_cross, b2b_total, b2b_cross) sheet names."""
    return (find_sheet(prefix, "B2C", "total"), find_sheet(prefix, "B2C", "cross"),
            find_sheet(prefix, "B2B", "total"), find_sheet(prefix, "B2B", "cross"))

datasets = {}
def build(key, label, prefix, qset):
    a, b, c, d = sheets_for(prefix)
    datasets[key] = {
        "label": label,
        "questionSet": qset,
        "employee": {"base": parse_total(wb[a]), "cross": parse_cross(wb[b])},
        "employer": {"base": parse_total(wb[c]), "cross": parse_cross(wb[d])},
    }

build(TOTAL_KEY, TOTAL_LABEL, None, "Általános")
for prefix, name in AREA_SHEETS.items():
    build(name, name, prefix, "IT + Contracting" if prefix in IT_SET else "Általános")

# Re-apply the client's approved editorial changes from the previous round — the workbook's
# topic sheet is the raw source and predates them.
TOPIC_RENAME = {
    "Munkahely váltás": "Munkahely váltás és toborzási kilátások",
    "Béremelés és juttatások": "Bérezés és juttatások",
}
EU_Q_PREFIX = "Szükséges az EU bértranszparencia"
mv_key = "Munkahely váltás"
ber_key = "Béremelés és juttatások"
for qs in ("Általános", "IT + Contracting"):
    mv = topic_data[mv_key][qs]["Munkáltatók"]
    eu = [q for q in mv if q.startswith(EU_Q_PREFIX)]
    for q in eu:
        mv.remove(q)
        topic_data[ber_key][qs]["Munkáltatók"].append(q)
    if not eu:
        sys.exit(f"EU transparency question not found in {qs}")

topics = [{"topic": TOPIC_RENAME.get(t, t), "questionSets": topic_data[t]} for t in topics_order]

# Validation: every question a topic lists must have base data in every dataset using that set.
problems = []
for t in topics:
    for qs, sides in t["questionSets"].items():
        for side_key, side in (("Munkavállalók", "employee"), ("Munkáltatók", "employer")):
            for q in sides[side_key]:
                for dk, d in datasets.items():
                    if d["questionSet"] != qs:
                        continue
                    if q not in d[side]["base"]:
                        problems.append(f"{d['label']} / {side} / {t['topic']}: MISSING {q[:60]}")
if problems:
    print(f"!! {len(problems)} missing question(s):")
    for x in problems[:12]:
        print("   ", x)
else:
    print("validation: every topic question resolves in every dataset")

# preserve everything that isn't survey data
old = json.load(io.open(OUT, encoding="utf-8"))
new = {
    "generatedFrom": "Salary Guide- kutatási eredmények-final.xlsx",
    "totalKey": TOTAL_KEY,
    "totalLabel": TOTAL_LABEL,
    "topics": topics,
    "datasets": datasets,
    "salary": old["salary"],
    "sapProducts": old["sapProducts"],
    "siteMap": old["siteMap"],
    "filterDimensions": old["filterDimensions"],
}

# The salary block (including the Talent Insight counts) is owned by build-salary-data.py.
io.open(OUT, "w", encoding="utf-8").write(json.dumps(new, ensure_ascii=False, separators=(",", ":")))

# ---- report ----
print(f"datasets: {len(datasets)}")
for k, v in datasets.items():
    print(f"  {v['label']:34s} set={v['questionSet']:16s} "
          f"emp base/cross={len(v['employee']['base'])}/{len(v['employee']['cross'])} "
          f"empr base/cross={len(v['employer']['base'])}/{len(v['employer']['cross'])}")
print("\ntopics:")
for t in topics:
    for qs, sides in t["questionSets"].items():
        print(f"  {t['topic']:38s} {qs:16s} MV={len(sides['Munkavállalók'])} MT={len(sides['Munkáltatók'])}")
w = [q for d in datasets.values() for q in d["employee"]["base"]
     if d["employee"]["base"][q]["kind"] == "weighted"]
print(f"\nweighted-only question instances: {len(w)}")
print("sample:", w[0][:70] if w else "NONE")
